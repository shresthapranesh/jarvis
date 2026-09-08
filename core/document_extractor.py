import base64
import io
import re


MAX_CHARS = 80_000

# Tabular formats: files whose value is in their *structure*, not their prose.
# Flattening these to text and embedding the result produces retrieval that
# structurally cannot answer the question anyone actually asks of them — an
# aggregate, a count, a filter — so they are routed to the kernel by path
# instead of into the prompt. See core/streaming.py:_tabular_part.
_TABULAR_EXTS = (".csv", ".tsv", ".xlsx", ".xls", ".parquet", ".jsonl", ".ndjson")
_TABULAR_MIMES = {
    "text/csv",
    "text/tab-separated-values",
    "application/csv",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.apache.parquet",
    "application/x-parquet",
    "application/x-ndjson",
    "application/jsonl",
}
# The subset whose own bytes are line-oriented text, so a head-of-file read is
# the real content rather than a decode of a binary container.
_TEXT_TABULAR_EXTS = (".csv", ".tsv", ".jsonl", ".ndjson")
_TEXT_TABULAR_MIMES = {
    "text/csv",
    "text/tab-separated-values",
    "application/csv",
    "application/x-ndjson",
    "application/jsonl",
}


def is_tabular(mime_type: str, filename: str) -> bool:
    """True for row/column data the agent should open with code, not read as text."""
    return (
        mime_type.lower() in _TABULAR_MIMES
        or filename.lower().endswith(_TABULAR_EXTS)
    )


def is_text_tabular(mime_type: str, filename: str) -> bool:
    """True when the file's own bytes are line-oriented text, so a head read previews it.

    Narrower than `is_tabular` on purpose: previewing the first bytes of an
    .xlsx or .parquet yields container framing, not rows.
    """
    return (
        mime_type.lower() in _TEXT_TABULAR_MIMES
        or filename.lower().endswith(_TEXT_TABULAR_EXTS)
    )


def extract_raw_text(mime_type: str, base64_data: str, filename: str) -> str:
    """Extract the full, untruncated document text. Raises on failure.

    Used by the chunk indexer (core/doc_index.py), which needs the whole
    document; `extract_text` below is the inline-friendly wrapper.
    """
    raw = base64.b64decode(base64_data)
    mt = mime_type.lower()
    fn = filename.lower()

    if mt == "application/pdf" or fn.endswith(".pdf"):
        return _extract_pdf(raw)
    if mt in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ) or fn.endswith((".docx", ".doc")):
        return _extract_docx(raw)
    if mt in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ) or fn.endswith((".xlsx", ".xls")):
        return _extract_xlsx(raw)
    if mt.startswith("text/") or fn.endswith((".txt", ".md", ".csv", ".tsv")):
        return raw.decode("utf-8", errors="replace")
    if mt in ("application/rtf", "text/rtf") or fn.endswith(".rtf"):
        return _extract_rtf(raw)
    return raw.decode("utf-8", errors="replace")


def format_inline(filename: str, text: str) -> str:
    """Wrap extracted text in the inline document framing, capped at MAX_CHARS."""
    header = f"[Document: {filename}]\n"
    footer = "\n[End of document]"
    if len(text) > MAX_CHARS:
        body = text[:MAX_CHARS] + f"\n... [truncated — {len(text) - MAX_CHARS} chars omitted]"
    else:
        body = text
    return header + body + footer


def extract_text(mime_type: str, base64_data: str, filename: str) -> str:
    try:
        return format_inline(filename, extract_raw_text(mime_type, base64_data, filename))
    except Exception as e:
        return f"[Document: {filename}]\n[Extraction failed: {e}]\n[End of document]"


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(raw: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _extract_xlsx(raw: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheets = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append("\t".join(cells))
        if rows:
            sheets.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(sheets)


def _extract_rtf(raw: bytes) -> str:
    text = raw.decode("latin-1", errors="replace")
    text = re.sub(r"\{[^{}]*\}|\\[a-z]+\d* ?|[{}]", "", text)
    return text.strip()
